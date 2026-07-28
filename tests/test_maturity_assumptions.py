from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from src.asx_maturity_screen import (
    DebtAgreement, DebtTranche, ExtractionPayload, MoneyValue,
    committed_agreement_totals, derive_tranche_screening, maturity_grid,
    nearest_screening_maturity, parse_tenor_months, postprocess_agreements,
    split_multi_tenor_tranches, validate_extraction, workbook,
)


def money(value):
    return None if value is None else MoneyValue(value_m=value, currency="AUD", source_page=10, source_quote_or_evidence="evidence")


def tranche(name="Facility A", description=None, exact=None, drawn=None, limit=None,
            reference=None, reference_type="UNDETERMINED", tenor=None, tenor_text=None,
            status="UNAVAILABLE", evidence=None, currency="AUD", instrument="REVOLVING_FACILITY"):
    return DebtTranche(
        parent_agreement_id="", ticker="TST", tranche_name=name,
        tranche_description=description, tranche_instrument_type=instrument,
        tranche_currency=currency, tranche_limit=money(limit), tranche_drawn_amount=money(drawn),
        tranche_undrawn_amount=None, maturity_description=description,
        maturity_reference_date=reference, maturity_reference_type=reference_type,
        tenor_months=tenor, tenor_description=tenor_text, tenor_basis="DISCLOSED_REMAINING_TENOR" if tenor or tenor_text else "UNDETERMINED",
        exact_maturity_date=exact, derived_maturity_date=None,
        assumed_earliest_maturity_date=None, screening_maturity_date=None,
        screening_maturity_is_assumed=False, maturity_assumption_basis="UNDETERMINED",
        maturity_assumption_explanation=None, screening_amount_m=None,
        screening_amount_basis="UNAVAILABLE", amount_allocation_status=status,
        source_page=10, source_quote_or_evidence=evidence or description, confidence=0.9,
    )


def agreement(tranches, name="Club debt agreement", committed="COMMITTED",
              limit=468.9, drawn=341.252, refinancing="2024-12-20", current=True):
    return DebtAgreement(
        agreement_id="model-id", ticker="TST", agreement_name=name,
        agreement_description="Refinanced club debt agreement", lender_or_market="Club lenders",
        instrument_type="CLUB_FACILITY", currency="AUD", secured_or_unsecured="UNSECURED",
        agreement_facility_limit=money(limit), agreement_drawn_amount=money(drawn),
        agreement_undrawn_amount=None, committed_or_uncommitted=committed,
        refinancing_date=refinancing, refinancing_date_basis="DISCLOSED_REFINANCING_DATE",
        effective_date=None, effective_date_basis=None, financial_close_date=None,
        financial_close_date_basis=None, agreement_source_page=10,
        agreement_source_quote_or_evidence="Current club facility disclosure",
        confidence=0.9, is_current_period=current, tranches=tranches,
    )


def payload(agreements, balance="2025-06-30"):
    return ExtractionPayload(
        ticker="TST", company_name="Test Issuer", matched_legal_entity="Test Issuer Limited",
        financial_year="2025", balance_date=balance, reporting_currency="AUD", reporting_unit="$m",
        gross_debt_excluding_leases=money(341.252), current_borrowings_excluding_leases=money(0),
        non_current_borrowings_excluding_leases=money(341.252), undrawn_committed_headroom=money(127.648),
        cash_and_cash_equivalents=money(20), net_debt_excluding_leases=money(321.252),
        lease_liabilities=money(5), debt_agreements=agreements, disclosure_buckets=[],
        extraction_confidence=0.9, extraction_status="EXTRACTED", model_used="mock",
        source_pages=[10], extraction_notes="", validation_flags=[], profile_quality=None,
        leases_apparently_included=False,
    )


def test_written_numeric_decimal_and_combined_tenors():
    assert parse_tenor_months("two years") == 24
    assert parse_tenor_months("2 years, 6 months") == 30
    assert parse_tenor_months("four years, six months") == 54
    assert parse_tenor_months("4.5 years") == 54
    assert parse_tenor_months("30 months") == 30
    assert parse_tenor_months("five-year revolving facility") == 60


def test_multi_tenor_sentence_splits_named_facilities():
    evidence = "two years, six months (Facility A) and four years, six months (Facility B)"
    rows = split_multi_tenor_tranches(agreement([tranche(name="Combined", evidence=evidence)]))
    assert [(row.tranche_name, row.tenor_months) for row in rows] == [("Facility A", 30), ("Facility B", 54)]


def test_reference_tenor_derivation_and_month_start():
    for reference_type, reference in (
        ("BALANCE_DATE", "2025-06-30"), ("REFINANCING_DATE", "2024-12-20"),
        ("EFFECTIVE_DATE", "2025-01-15"), ("FINANCIAL_CLOSE_DATE", "2025-02-28"),
    ):
        item = tranche(reference=reference, reference_type=reference_type, tenor=30)
        result = derive_tranche_screening(item, "2025-06-30", agreement([item]))
        expected = date.fromisoformat(reference)
        assert result.exact_maturity_date is None
        assert result.derived_maturity_date is not None
        assert result.screening_maturity_date.endswith("-01")
        assert result.maturity_assumption_basis == "RELATIVE_TENOR_MONTH_START"
        assert result.screening_maturity_date != expected.isoformat()
    wording = tranche(
        description="remaining tenor of two years, six months measured from 30 June 2025",
        reference_type="UNDETERMINED",
    )
    wording = derive_tranche_screening(wording, "2025-06-30", agreement([wording]))
    assert wording.maturity_reference_date == "2025-06-30"
    assert wording.exact_maturity_date is None
    assert wording.derived_maturity_date == "2027-12-30"
    assert wording.screening_maturity_date == "2027-12-01"
    refinanced = tranche(reference_type="REFINANCING_DATE", tenor=30)
    refinanced = derive_tranche_screening(refinanced, "2025-06-30", agreement([refinanced]))
    assert refinanced.maturity_reference_date == "2024-12-20"
    assert refinanced.derived_maturity_date == "2027-06-20"
    assert refinanced.screening_maturity_date == "2027-06-01"


def test_representative_two_tranche_agreement_and_unallocated_amounts():
    a = tranche("Facility A", reference_type="BALANCE_DATE", tenor=30, status="AGREEMENT_LEVEL_ONLY")
    b = tranche("Facility B", reference_type="BALANCE_DATE", tenor=54, status="AGREEMENT_LEVEL_ONLY")
    uncommitted = agreement([], name="Uncommitted accordion", committed="UNCOMMITTED", limit=200, drawn=None)
    result = payload([agreement([a, b]), uncommitted])
    postprocess_agreements(result)
    current = next(row for row in result.debt_agreements if row.committed_or_uncommitted == "COMMITTED")
    assert money_value(current.agreement_undrawn_amount) == 127.648
    assert current.refinancing_date == "2024-12-20"
    assert [(row.tranche_name, row.tenor_months, row.derived_maturity_date, row.screening_maturity_date) for row in current.tranches] == [
        ("Facility A", 30, "2027-12-30", "2027-12-01"),
        ("Facility B", 54, "2029-12-30", "2029-12-01"),
    ]
    for row in current.tranches:
        assert row.screening_amount_m is None
        assert row.screening_amount_basis == "AGREEMENT_AMOUNT_UNALLOCATED"
        assert row.amount_allocation_status == "AGREEMENT_LEVEL_ONLY"
    grid, _ = maturity_grid(result)
    assert all(value is None for value in grid.values())
    extraction = result.model_dump(mode="json")
    nearest = nearest_screening_maturity(extraction)
    assert nearest["Nearest tranche name"] == "Facility A"
    assert nearest["Nearest tranche maturity date"] == "2027-12-01"
    assert nearest["Amount potentially maturing"] is None
    assert committed_agreement_totals(extraction) == (468.9, 127.648)
    assert "TRANCHE_AMOUNT_UNDISCLOSED" in validate_extraction(result, "2025")


def money_value(value):
    return value.value_m if value else None


def test_tranche_amount_priority_partial_and_grid_separation():
    exact = tranche("Bond series", exact="2027-12-15", drawn=50, limit=60,
                    status="TRANCHE_LEVEL_DISCLOSED", instrument="BOND")
    assumed = tranche("Private placement", description="December 2027", drawn=100, limit=120,
                      status="TRANCHE_LEVEL_DISCLOSED", instrument="PRIVATE_PLACEMENT")
    partial = tranche("USD term tranche", description="2028", drawn=25, limit=None,
                      status="PARTIAL_TRANCHE_ALLOCATION", currency="USD", instrument="PRIVATE_PLACEMENT")
    result = payload([agreement([exact, assumed, partial])])
    postprocess_agreements(result)
    rows = result.debt_agreements[0].tranches
    assert rows[0].screening_amount_basis == "TRANCHE_DRAWN_AMOUNT"
    assert rows[2].amount_allocation_status == "PARTIAL_TRANCHE_ALLOCATION"
    assert {row.tranche_currency for row in rows} == {"AUD", "USD"}
    assert {row.tranche_instrument_type for row in rows} == {"BOND", "PRIVATE_PLACEMENT"}
    grid, _ = maturity_grid(result)
    assert grid["2H27 exact"] == 50
    assert grid["2H27 inferred"] == 100
    assert grid["2H27 total"] == 150


def test_current_period_filter_and_stable_ids():
    current = agreement([tranche()], current=True)
    prior = agreement([tranche(name="Prior year tranche")], name="Prior year agreement", current=False)
    first = payload([current, prior])
    second = payload([current.model_copy(deep=True)])
    postprocess_agreements(first)
    postprocess_agreements(second)
    assert len(first.debt_agreements) == 1
    assert "PRIOR_PERIOD_DATA_USED" in first.validation_flags
    assert first.debt_agreements[0].agreement_id == second.debt_agreements[0].agreement_id


def test_workbook_agreement_tranche_columns_dates_and_hyperlink(tmp_path: Path):
    result = payload([agreement([tranche(exact="2027-12-15", drawn=50, limit=60)])])
    postprocess_agreements(result)
    grid, quality = maturity_grid(result)
    selected = tmp_path / "report.pdf"
    selected.write_bytes(b"pdf")
    match = {"ticker": "TST", "target_name": "Test Issuer", "match_status": "MATCHED_HIGH", "match_confidence": 100,
             "selected_pdf": str(selected), "report_type": "ANNUAL_REPORT", "financial_year": "2025"}
    record = {"ticker": "TST", "target_name": "Test Issuer", "document_match": match,
              "extraction": result.model_dump(mode="json"), "status": "EXTRACTED", "validation_flags": [],
              "maturity_grid": grid, "profile_quality": quality, "selected_pages": [10], "run_log": {}, "error": ""}
    output = tmp_path / "screen.xlsx"
    workbook(output, [record], [match])
    book = load_workbook(output)
    sheet = book["Facilities & Instruments"]
    headers = [cell.value for cell in sheet[1]]
    assert headers[:5] == ["ticker", "company_name", "agreement_id", "agreement_name", "tranche_name"]
    date_cell = sheet.cell(2, headers.index("screening_maturity_date") + 1)
    assert date_cell.number_format == "dd/mm/yyyy"
    assert sheet.cell(2, headers.index("selected_pdf") + 1).hyperlink is not None
    summary_headers = [cell.value for cell in book["Summary"][1]]
    link = book["Summary"].cell(2, summary_headers.index("Selected PDF") + 1)
    assert link.hyperlink is not None
