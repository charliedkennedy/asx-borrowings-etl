from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from src.asx_maturity_screen import (
    DebtAgreement, DebtTranche, DisclosureBucket, ExtractionPayload, MoneyValue,
    derive_tranche_screening, maturity_grid, nearest_screening_maturity,
    postprocess_agreements, remaining_annual_instalment_dates, summary_row, validate_extraction, workbook,
)


def cash(value, currency="AUD", role="UNKNOWN"):
    return MoneyValue(value_m=value, currency=currency, source_page=1,
                      source_quote_or_evidence="fixture evidence", amount_role=role)


def tranche(name="Facility A", **changes):
    values = dict(
        parent_agreement_id="", ticker="TST", tranche_name=name, tranche_description=None,
        tranche_instrument_type="REVOLVING", tranche_currency="AUD", tranche_limit=None,
        tranche_drawn_amount=None, tranche_undrawn_amount=None, maturity_description=None,
        maturity_reference_date=None, maturity_reference_type="UNDETERMINED", tenor_months=None,
        tenor_description=None, tenor_basis="UNDETERMINED", exact_maturity_date=None,
        derived_maturity_date=None, assumed_earliest_maturity_date=None,
        screening_maturity_date=None, screening_maturity_is_assumed=False,
        maturity_assumption_basis="UNDETERMINED", maturity_assumption_explanation=None,
        screening_amount_m=None, screening_amount_basis="UNAVAILABLE",
        amount_allocation_status="UNAVAILABLE", source_page=1,
        source_quote_or_evidence="fixture evidence", confidence=.9,
    )
    values.update(changes)
    return DebtTranche(**values)


def agreement(children, **changes):
    values = dict(
        agreement_id="", ticker="TST", agreement_name="Test facility",
        agreement_description=None, lender_or_market=None, instrument_type="REVOLVING",
        currency="AUD", secured_or_unsecured=None, agreement_facility_limit=None,
        agreement_drawn_amount=None, agreement_undrawn_amount=None,
        committed_or_uncommitted="COMMITTED", refinancing_date=None,
        refinancing_date_basis=None, effective_date=None, effective_date_basis=None,
        financial_close_date=None, financial_close_date_basis=None, agreement_source_page=1,
        agreement_source_quote_or_evidence="Test facility fixture evidence", confidence=.9,
        is_current_period=True, tranches=children,
    )
    values.update(changes)
    return DebtAgreement(**values)


def payload(agreements=None, buckets=None):
    return ExtractionPayload(
        ticker="TST", company_name="Model name ignored", matched_legal_entity="Test Legal Limited",
        financial_year="2025", balance_date="2025-06-30", reporting_currency="AUD",
        reporting_unit="$m", gross_debt_excluding_leases=cash(100),
        current_borrowings_excluding_leases=cash(20), non_current_borrowings_excluding_leases=cash(80),
        undrawn_committed_headroom=None, cash_and_cash_equivalents=cash(120),
        net_debt_excluding_leases=cash(-20), lease_liabilities=None,
        debt_agreements=agreements or [], disclosure_buckets=buckets or [], extraction_confidence=.9,
        extraction_status="EXTRACTED", model_used="mock", source_pages=[1], extraction_notes="",
        validation_flags=[], profile_quality=None, leases_apparently_included=False,
    )


def test_as_of_date_excludes_historical_maturity_but_retains_detail():
    row = tranche(exact_maturity_date="2025-12-15", tranche_drawn_amount=cash(50, role="FACILITY_DRAWN"))
    data = payload([agreement([row])])
    postprocess_agreements(data)
    grid, _ = maturity_grid(data, date(2026, 7, 1))
    assert all(grid[f"{half} total"] is None for half in ("2H26", "1H27", "2H27", "1H28", "2H28", "FY29+"))
    assert nearest_screening_maturity(data.model_dump(mode="json"), date(2026, 7, 1)) == {}
    assert "MATURITY_BEFORE_SCREENING_AS_OF_DATE" in validate_extraction(data, "2025", date(2026, 7, 1))
    assert data.debt_agreements[0].tranches[0].screening_maturity_date == "2025-12-15"


def test_measurement_date_is_not_maturity_and_upper_bound_is_not_earliest():
    measurement = derive_tranche_screening(tranche(
        maturity_description="fully utilised at 30 June 2025",
        maturity_reference_date="2025-06-30", maturity_reference_type="BALANCE_DATE",
    ), "2025-06-30")
    assert measurement.screening_maturity_date is None
    upper = derive_tranche_screening(tranche(
        maturity_description="facility maturities staggered through to September 2027",
    ), "2025-06-30")
    assert upper.screening_maturity_date is None
    assert "latest/final maturity endpoint" in upper.maturity_assumption_explanation


def test_financial_year_period_uses_actual_year_end():
    june = derive_tranche_screening(tranche(maturity_description="Financial Year of Debt Maturity FY2026"), "2025-06-30")
    december = derive_tranche_screening(tranche(maturity_description="Financial Year of Debt Maturity FY2026"), "2025-12-31")
    assert (june.screening_maturity_date, june.maturity_period_end) == ("2025-07-01", "2026-06-30")
    assert (december.screening_maturity_date, december.maturity_period_end) == ("2026-01-01", "2026-12-31")
    assert june.maturity_assumption_basis == "FINANCIAL_YEAR_START_ASSUMPTION"


def test_contractual_buckets_get_boundaries_and_do_not_become_debt_grid():
    buckets = [DisclosureBucket(
        ticker="TST", bucket_label="between one and two years", period_start=None, period_end=None,
        amount=cash(40, role="CONTRACTUAL_CASH_FLOW"), currency="AUD", source_page=1,
        source_quote_or_evidence="contractual cash flows including interest", confidence=.8,
        amount_type="PRINCIPAL_AND_INTEREST", safe_for_summary=False,
    )]
    data = payload(buckets=buckets)
    postprocess_agreements(data)
    assert (buckets[0].period_start, buckets[0].period_end) == ("2026-07-01", "2027-06-30")
    grid, quality = maturity_grid(data)
    assert all(grid[f"{half} total"] is None for half in ("2H26", "1H27", "2H27", "1H28", "2H28", "FY29+"))
    assert quality == "NO_MATURITY_DISCLOSURE"


def test_only_remaining_unambiguous_annual_instalments_are_derived():
    text = ("Four equal annual payments. The second payment was made on 28 February 2025 "
            "and two instalments remaining.")
    assert [item.isoformat() for item in remaining_annual_instalment_dates(text)] == [
        "2026-02-28", "2027-02-28",
    ]
    assert remaining_annual_instalment_dates("annual payments may be made") == []


def test_foreign_currency_and_facility_limit_are_blocked_from_debt_grid_but_capacity_is_separate():
    row = tranche(
        exact_maturity_date="2027-11-15", tranche_limit=cash(400, "AUD", "FACILITY_LIMIT"),
        tranche_drawn_amount=None, exposure_type="FACILITY_CAPACITY",
    )
    foreign = tranche(
        "USD note", exact_maturity_date="2027-11-20", tranche_drawn_amount=cash(125, "USD", "PRINCIPAL_OUTSTANDING"),
        tranche_currency="USD", exposure_type="FUNDED_DEBT",
    )
    data = payload([agreement([row, foreign])])
    postprocess_agreements(data)
    grid, _ = maturity_grid(data)
    assert grid["2H27 total"] is None
    assert grid["2H27 capacity total"] == 400
    assert "REPORTING_CURRENCY_EQUIVALENT_UNAVAILABLE" in validate_extraction(data, "2025")


def test_single_child_amount_propagates_but_multiple_children_stay_unallocated():
    only = tranche(source_quote_or_evidence="Test facility drawn amount and limit")
    one = payload([agreement([only], agreement_drawn_amount=cash(30), agreement_facility_limit=cash(50))])
    postprocess_agreements(one)
    assert one.debt_agreements[0].tranches[0].tranche_drawn_amount.value_m == 30
    many = payload([agreement([tranche(), tranche("Facility B")], agreement_drawn_amount=cash(30))])
    postprocess_agreements(many)
    assert all(item.tranche_drawn_amount is None for item in many.debt_agreements[0].tranches)


def test_cross_currency_or_note_basis_never_calculates_negative_undrawn():
    note = agreement(
        [tranche("USPP series", face_value=cash(300, "USD", "FACE_VALUE"),
                 carrying_amount=cash(457.364, "AUD", "CARRYING_AMOUNT"), tranche_currency="USD")],
        instrument_type="USPP", agreement_facility_limit=cash(300, "USD", "FACE_VALUE"),
        agreement_drawn_amount=cash(457.364, "AUD", "CARRYING_AMOUNT"),
    )
    data = payload([note])
    postprocess_agreements(data)
    assert data.debt_agreements[0].agreement_undrawn_amount is None
    assert "CURRENCY_OR_BASIS_MISMATCH" in validate_extraction(data, "2025")


def test_workbook_has_review_queue_target_company_and_net_debt_fields(tmp_path: Path):
    data = payload()
    match = {"ticker": "TST", "target_name": "Target Company Limited", "match_status": "MATCHED_HIGH",
             "match_confidence": 100, "selected_pdf": "", "report_type": "ANNUAL_REPORT"}
    record = {"ticker": "TST", "target_name": "Target Company Limited", "document_match": match,
              "extraction": data.model_dump(mode="json"), "status": "EXTRACTED", "validation_flags": [],
              "maturity_grid": {}, "profile_quality": "NO_MATURITY_DISCLOSURE", "selected_pages": [1],
              "error": "", "run_log": {}, "screening_as_of_date": "2026-07-29"}
    output = tmp_path / "screen.xlsx"
    workbook(output, [record], [match])
    book = load_workbook(output)
    assert book.sheetnames[-1] == "Review Queue"
    headers = [cell.value for cell in book["Summary"][1]]
    values = [cell.value for cell in book["Summary"][2]]
    row = dict(zip(headers, values))
    assert row["Company name"] == "Target Company Limited"
    assert row["Calculated net debt ex leases"] == -20
    assert row["Screening as-of date"].strftime("%Y-%m-%d") == "2026-07-29"
