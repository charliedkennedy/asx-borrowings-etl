import csv
from pathlib import Path

import fitz
import pytest
from openpyxl import load_workbook

from src.asx_maturity_screen import ExtractionPayload, main_for_args, parser


COHORT = {
    "IDX": "Integral Diagnostics",
    "ORA": "Orora Limited",
    "KPG": "Kelly Partners Group Holdings Limited",
    "CMW": "Cromwell Property Group",
    "ALQ": "ALS Limited",
    "WGN": "Wagners Holding Company Limited",
}


def make_pdf(path: Path, company: str) -> None:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), f"{company}\nAnnual Report\nBorrowings and financing facilities")
    document.save(path)
    document.close()


def setup_files(tmp_path: Path) -> tuple[Path, Path]:
    root = tmp_path / "local reports"
    root.mkdir()
    rows = ["ticker,target_name,aliases"]
    for ticker, company in {**COHORT, "MTS": "Metcash Limited"}.items():
        make_pdf(root / f"{ticker}_2025-08-25_FY25_Annual_Report.pdf", company)
        rows.append(f"{ticker},{company},{company}")
    targets = tmp_path / "targets.csv"
    targets.write_text("\n".join(rows) + "\n", encoding="utf-8")
    return root, targets


def extracted(ticker: str) -> ExtractionPayload:
    return ExtractionPayload(
        ticker=ticker, company_name=COHORT.get(ticker, ticker), matched_legal_entity=None,
        financial_year="2025", balance_date="2025-06-30", reporting_currency="AUD",
        reporting_unit="$m", gross_debt_excluding_leases=None,
        current_borrowings_excluding_leases=None, non_current_borrowings_excluding_leases=None,
        undrawn_committed_headroom=None, cash_and_cash_equivalents=None,
        net_debt_excluding_leases=None, lease_liabilities=None, debt_agreements=[],
        disclosure_buckets=[], extraction_confidence=0.9, extraction_status="EXTRACTED",
        model_used="mock", source_pages=[1], extraction_notes="No debt disclosed in fixture.",
        validation_flags=[], profile_quality="SPLIT_ONLY", leases_apparently_included=False,
    )


def test_multi_ticker_combined_workbook_resumability_and_force(monkeypatch, tmp_path: Path) -> None:
    root, targets = setup_files(tmp_path)
    output = tmp_path / "outputs" / "screen.xlsx"
    monkeypatch.setenv("OPENAI_API_KEY", "mock-key")
    calls: list[str] = []

    def extractor(ticker, *args):
        calls.append(ticker)
        return extracted(ticker), {"input_tokens": 1, "output_tokens": 1}

    requested = ["idx", "ORA", "kpg", "CMW", "alq", "WGN"]
    command = ["--tickers", *requested, "--pdf-root", str(root), "--targets", str(targets), "--output", str(output)]
    assert main_for_args(command, extractor=extractor) == 0
    assert calls == list(COHORT)
    book = load_workbook(output)
    summary = book["Summary"]
    assert {summary.cell(row, 1).value for row in range(2, summary.max_row + 1)} == set(COHORT)
    assert "MTS" not in {summary.cell(row, 1).value for row in range(2, summary.max_row + 1)}

    assert main_for_args(command, extractor=extractor) == 0
    assert calls == list(COHORT)
    assert main_for_args([*command, "--force"], extractor=extractor) == 0
    assert calls == [*COHORT, *COHORT]


def test_multi_ticker_match_only_is_key_free_and_exact(monkeypatch, tmp_path: Path) -> None:
    root, targets = setup_files(tmp_path)
    output = tmp_path / "outputs" / "match.xlsx"
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    calls = 0

    def forbidden(*args):
        nonlocal calls
        calls += 1
        raise AssertionError("matching-only must not call extraction")

    command = ["--tickers", *COHORT, "--match-only", "--pdf-root", str(root),
               "--targets", str(targets), "--output", str(output)]
    assert main_for_args(command, extractor=forbidden) == 0
    assert calls == 0
    with (output.parent / "document_register.csv").open(encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert [row["ticker"] for row in rows] == list(COHORT)
    assert all(row["match_status"] == "MATCHED_HIGH" for row in rows)


def test_missing_ticker_and_selector_conflicts_are_clear(tmp_path: Path) -> None:
    root, targets = setup_files(tmp_path)
    with pytest.raises(SystemExit, match="ZZZ"):
        main_for_args(["--tickers", "IDX", "ZZZ", "--match-only", "--pdf-root", str(root), "--targets", str(targets)])
    base = ["--pdf-root", str(root), "--targets", str(targets)]
    with pytest.raises(SystemExit):
        parser().parse_args([*base, "--ticker", "IDX", "--tickers", "ORA"])
    with pytest.raises(SystemExit):
        parser().parse_args([*base, "--pilot", "--tickers", "IDX"])


def test_multi_ticker_extraction_stops_unless_every_match_is_high(monkeypatch, tmp_path: Path) -> None:
    root, targets = setup_files(tmp_path)
    duplicate_dir = root / "duplicate"
    duplicate_dir.mkdir()
    source = root / "ORA_2025-08-25_FY25_Annual_Report.pdf"
    (duplicate_dir / source.name).write_bytes(source.read_bytes())
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    calls = 0

    def forbidden(*args):
        nonlocal calls
        calls += 1

    assert main_for_args(["--tickers", *COHORT, "--pdf-root", str(root),
                          "--targets", str(targets), "--output", str(tmp_path / "out" / "screen.xlsx")],
                         extractor=forbidden) == 3
    assert calls == 0


def test_existing_single_ticker_and_pilot_selection(monkeypatch, tmp_path: Path) -> None:
    root, targets = setup_files(tmp_path)
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    single_output = tmp_path / "single" / "screen.xlsx"
    assert main_for_args(["--ticker", "ORA", "--match-only", "--pdf-root", str(root),
                          "--targets", str(targets), "--output", str(single_output)]) == 0
    with (single_output.parent / "document_register.csv").open(encoding="utf-8", newline="") as handle:
        assert [row["ticker"] for row in csv.DictReader(handle)] == ["ORA"]
    pilot_output = tmp_path / "pilot" / "screen.xlsx"
    assert main_for_args(["--pilot", "--match-only", "--pdf-root", str(root),
                          "--targets", str(targets), "--output", str(pilot_output)]) == 0
    with (pilot_output.parent / "document_register.csv").open(encoding="utf-8", newline="") as handle:
        tickers = {row["ticker"] for row in csv.DictReader(handle)}
    assert tickers == {"IDX", "MTS"}


def test_repository_targets_contains_six_name_cohort() -> None:
    with Path("targets.csv").open(encoding="utf-8", newline="") as handle:
        rows = {row["ticker"]: row["target_name"] for row in csv.DictReader(handle)}
    assert {ticker: rows[ticker] for ticker in COHORT} == COHORT
